from rest_framework.decorators import api_view
from rest_framework.response import Response
from rest_framework import status
from django.utils import timezone
from decimal import Decimal
from .models import MaterialRequisition
import datetime
import json
from pymongo import MongoClient
from django.conf import settings

def get_mongo_collection(collection_name='franchise_materialrequisition'):
    try:
        client_dict = settings.DATABASES['default']['CLIENT']
        host = client_dict.get('host', 'mongodb://localhost:27017/')
        db_name = settings.DATABASES['default'].get('NAME', 'franchise')
        client = MongoClient(host)
        return client[db_name][collection_name]
    except Exception as e:
        print("Mongo connection error:", e)
        return None

def get_franchise_name_map():
    name_map = {}
    try:
        col_franchise = get_mongo_collection('franchise_franchise')
        col_location = get_mongo_collection('franchise_location_details')
        if col_franchise is not None:
            loc_map = {}
            if col_location is not None:
                for loc in col_location.find():
                    loc_id = loc.get('location_id')
                    cname = loc.get('Cluster_Name') or loc.get('District') or ''
                    if loc_id:
                        loc_map[loc_id] = cname.strip()
            
            for fran in col_franchise.find():
                fid = fran.get('franchise_id')
                fname = fran.get('franchise_name') or fid
                cluster = loc_map.get(fran.get('location_id'), '')
                display_name = f"{fname} ({cluster})".strip() if cluster else fname
                if fid:
                    name_map[fid] = display_name
    except Exception as e:
        print("Error getting franchise name map:", e)
    return name_map

def sync_doc_to_mongo(mr_number, update_dict):
    col = get_mongo_collection('franchise_materialrequisition')
    if col is not None:
        try:
            col.update_many({'mr_number': mr_number}, {'$set': update_dict})
        except Exception as e:
            print("Error syncing to Mongo:", e)

def generate_mr_number():
    now = datetime.datetime.now()
    year = now.year % 100
    next_year = (now.year + 1) % 100
    year_code = f"{year:02d}{next_year:02d}"
    prefix = f"MR/{year_code}/"
    
    col = get_mongo_collection('franchise_materialrequisition')
    max_seq = 0
    if col is not None:
        try:
            regex_pattern = f"^{prefix}"
            cursor = col.find({'mr_number': {'$regex': regex_pattern}}, {'mr_number': 1})
            for doc in cursor:
                mr = doc.get('mr_number', '')
                try:
                    seq = int(mr.split('/')[-1])
                    if seq > max_seq:
                        max_seq = seq
                except (ValueError, IndexError):
                    pass
        except Exception as e:
            print("Error finding max MR sequence in Mongo:", e)
            
    if max_seq == 0:
        last_req = MaterialRequisition.objects.filter(mr_number__startswith=prefix).order_by('-mr_number').first()
        if last_req:
            try:
                max_seq = int(last_req.mr_number.split('/')[-1])
            except ValueError:
                max_seq = 0
                
    new_seq = max_seq + 1
    return f"{prefix}{new_seq:06d}"

def format_date(d):
    if not d:
        return None
    if isinstance(d, (datetime.datetime, datetime.date)):
        return d.isoformat()
    return str(d)

def to_float(val):
    if val is None:
        return 0.0
    return float(str(val))

def parse_items(items):
    if not items:
        return []
    if isinstance(items, str):
        try:
            parsed = json.loads(items)
            return parsed if isinstance(parsed, list) else [parsed]
        except Exception:
            return []
    if isinstance(items, list):
        return items
    return [items]

@api_view(['GET', 'POST'])
def material_requisitions(request):
    if request.method == 'GET':
        franchise_id = request.query_params.get('franchise_id')
        status_filter = request.query_params.get('status')
        include_cancelled = request.query_params.get('include_cancelled', 'false').lower() == 'true'
        from_date = request.query_params.get('from_date')
        to_date = request.query_params.get('to_date')
        
        col = get_mongo_collection('franchise_materialrequisition')
        mongo_query = {}
        
        if not include_cancelled:
            mongo_query['$or'] = [
                {'is_cancelled': False},
                {'is_cancelled': {'$exists': False}},
                {'is_cancelled': None}
            ]
            
        if franchise_id:
            mongo_query['franchise_id'] = franchise_id
            
        if status_filter:
            mongo_query['status'] = status_filter

        if from_date or to_date:
            date_filter = {}
            if from_date:
                try:
                    f_dt = datetime.datetime.strptime(from_date, '%Y-%m-%d')
                    date_filter['$gte'] = f_dt
                except Exception:
                    pass
            if to_date:
                try:
                    t_dt = datetime.datetime.strptime(to_date, '%Y-%m-%d') + datetime.timedelta(days=1, microseconds=-1)
                    date_filter['$lte'] = t_dt
                except Exception:
                    pass
            if date_filter:
                mongo_query['created_date'] = date_filter
            
        docs = list(col.find(mongo_query).sort('created_date', -1)) if col is not None else []
        franchise_name_map = get_franchise_name_map()
        
        data = []
        for req in docs:
            fid = req.get('franchise_id') or ''
            fname = franchise_name_map.get(fid, fid)
            data.append({
                "mr_number": req.get('mr_number'),
                "franchise_id": fid,
                "franchise_name": fname,
                "items": parse_items(req.get('items')),
                "Total_amount": to_float(req.get('Total_amount')),
                "status": req.get('status', 'Draft'),
                "created_by": req.get('created_by', 'system'),
                "created_date": format_date(req.get('created_date')),
                "lastmodified_by": req.get('lastmodified_by', 'system'),
                "lastmodified_date": format_date(req.get('lastmodified_date')),
                "approved_by": req.get('approved_by', ''),
                "approved_date": format_date(req.get('approved_date')),
                "edited_by": req.get('edited_by', ''),
                "edited_reason": req.get('edited_reason', ''),
                "edited_date": format_date(req.get('edited_date')),
                "is_cancelled": bool(req.get('is_cancelled', False)),
                "is_cancelled_by": req.get('is_cancelled_by', ''),
                "is_cancelled_date": format_date(req.get('is_cancelled_date')),
            })
        return Response(data, status=status.HTTP_200_OK)

    elif request.method == 'POST':
        payload = request.data
        raw_items = payload.get('items', [])
        items = parse_items(raw_items)
        
        total_amount = Decimal('0.00')
        formatted_items = []
        for index, item in enumerate(items, start=1):
            if isinstance(item, dict):
                qty = int(item.get('quantity', 1))
                amt = float(item.get('amount', 0))
                name = item.get('item_name') or item.get('medicine_name') or ''
                formatted_item = {
                    "item_id": item.get('item_id', index),
                    "item_name": name,
                    "quantity": qty,
                    "amount": amt
                }
                formatted_items.append(formatted_item)
                total_amount += Decimal(str(amt)) * qty

        mr_num = payload.get('mr_number')
        if not mr_num:
            mr_num = generate_mr_number()

        user_name = payload.get('created_by', 'system')
        franchise_id = payload.get('franchise_id', 'DEFAULT')
        req_status = 'Draft'

        req = MaterialRequisition.objects.create(
            mr_number=mr_num,
            franchise_id=franchise_id,
            items=formatted_items,
            Total_amount=payload.get('Total_amount', total_amount),
            status=req_status,
            created_by=user_name,
            lastmodified_by=user_name,
        )

        sync_doc_to_mongo(req.mr_number, {'items': formatted_items})
        franchise_name_map = get_franchise_name_map()
        fname = franchise_name_map.get(franchise_id, franchise_id)

        return Response({
            "message": "Material Requisition saved successfully",
            "mr_number": req.mr_number,
            "data": {
                "mr_number": req.mr_number,
                "franchise_id": req.franchise_id,
                "franchise_name": fname,
                "items": formatted_items,
                "Total_amount": to_float(req.Total_amount),
                "status": req.status,
                "created_by": req.created_by,
                "created_date": format_date(req.created_date),
                "lastmodified_by": req.lastmodified_by,
                "lastmodified_date": format_date(req.lastmodified_date),
                "approved_by": req.approved_by,
                "approved_date": req.approved_date,
                "edited_by": req.edited_by,
                "edited_reason": req.edited_reason,
                "edited_date": req.edited_date,
                "is_cancelled": req.is_cancelled,
                "is_cancelled_by": req.is_cancelled_by,
                "is_cancelled_date": req.is_cancelled_date
            }
        }, status=status.HTTP_201_CREATED)

@api_view(['GET', 'PUT', 'DELETE'])
def material_requisition_detail(request, mr_number):
    try:
        req = MaterialRequisition.objects.get(mr_number=mr_number)
    except MaterialRequisition.DoesNotExist:
        return Response({"error": "Requisition not found"}, status=status.HTTP_404_NOT_FOUND)

    franchise_name_map = get_franchise_name_map()
    fid = req.franchise_id or ''
    fname = franchise_name_map.get(fid, fid)

    if request.method == 'GET':
        return Response({
            "mr_number": req.mr_number,
            "franchise_id": fid,
            "franchise_name": fname,
            "items": parse_items(req.items),
            "Total_amount": to_float(req.Total_amount),
            "status": req.status,
            "created_by": req.created_by,
            "created_date": format_date(req.created_date),
            "lastmodified_by": req.lastmodified_by,
            "lastmodified_date": format_date(req.lastmodified_date),
            "approved_by": req.approved_by,
            "approved_date": format_date(req.approved_date),
            "edited_by": req.edited_by,
            "edited_reason": req.edited_reason,
            "edited_date": format_date(req.edited_date),
            "is_cancelled": req.is_cancelled,
            "is_cancelled_by": req.is_cancelled_by,
            "is_cancelled_date": format_date(req.is_cancelled_date),
        })

    elif request.method == 'PUT':
        if req.status == 'Approved':
            return Response(
                {"error": "Approved requisitions cannot be edited or modified"}, 
                status=status.HTTP_400_BAD_REQUEST
            )

        payload = request.data
        mongo_updates = {}

        if 'items' in payload:
            formatted_items = []
            total_amount = Decimal('0.00')
            for index, item in enumerate(payload['items'], start=1):
                qty = int(item.get('quantity', 1))
                amt = float(item.get('amount', 0))
                name = item.get('item_name') or item.get('medicine_name') or ''
                formatted_item = {
                    "item_id": item.get('item_id', index),
                    "item_name": name,
                    "quantity": qty,
                    "amount": amt
                }
                formatted_items.append(formatted_item)
                total_amount += Decimal(str(amt)) * qty

            req.items = formatted_items
            req.Total_amount = payload.get('Total_amount', total_amount)
            mongo_updates['items'] = formatted_items

        if 'status' in payload:
            req.status = payload['status']
            mongo_updates['status'] = payload['status']

        if 'approved_by' in payload:
            req.approved_by = payload['approved_by']
            req.approved_date = timezone.now()
            mongo_updates['approved_by'] = payload['approved_by']
            mongo_updates['approved_date'] = req.approved_date

        if 'edited_reason' in payload or 'edited_by' in payload:
            req.edited_by = payload.get('edited_by', req.franchise_id or 'system')
            req.edited_reason = payload.get('edited_reason', '')
            req.edited_date = timezone.now()
            mongo_updates['edited_by'] = req.edited_by
            mongo_updates['edited_reason'] = req.edited_reason
            mongo_updates['edited_date'] = req.edited_date

        user_name = payload.get('lastmodified_by', 'system')
        req.lastmodified_by = user_name
        if req.Total_amount is not None:
            req.Total_amount = Decimal(str(req.Total_amount))
        req.save()

        if mongo_updates:
            sync_doc_to_mongo(req.mr_number, mongo_updates)

        return Response({"message": "Requisition updated successfully"})

    elif request.method == 'DELETE':
        if req.status == 'Approved':
            return Response(
                {"error": "Approved requisitions cannot be deleted"}, 
                status=status.HTTP_400_BAD_REQUEST
            )

        payload = request.data or {}
        cancelled_by = payload.get('is_cancelled_by', req.franchise_id or 'system')
        
        # Perform Soft Delete
        req.is_cancelled = True
        req.is_cancelled_by = cancelled_by
        req.is_cancelled_date = timezone.now()
        if req.Total_amount is not None:
            req.Total_amount = Decimal(str(req.Total_amount))
        req.save()

        sync_doc_to_mongo(req.mr_number, {
            'is_cancelled': True,
            'is_cancelled_by': cancelled_by,
            'is_cancelled_date': req.is_cancelled_date
        })

        return Response({"message": "Requisition cancelled (soft deleted) successfully"}, status=status.HTTP_200_OK)

@api_view(['GET'])
def export_material_requisitions_excel(request):
    franchise_id = request.query_params.get('franchise_id')
    status_filter = request.query_params.get('status')
    from_date = request.query_params.get('from_date')
    to_date = request.query_params.get('to_date')
    
    col = get_mongo_collection('franchise_materialrequisition')
    mongo_query = {'$or': [{'is_cancelled': False}, {'is_cancelled': {'$exists': False}}, {'is_cancelled': None}]}
    
    if franchise_id:
        mongo_query['franchise_id'] = franchise_id
    if status_filter:
        mongo_query['status'] = status_filter
    if from_date or to_date:
        date_filter = {}
        if from_date:
            try:
                date_filter['$gte'] = datetime.datetime.strptime(from_date, '%Y-%m-%d')
            except Exception:
                pass
        if to_date:
            try:
                date_filter['$lte'] = datetime.datetime.strptime(to_date, '%Y-%m-%d') + datetime.timedelta(days=1, microseconds=-1)
            except Exception:
                pass
        if date_filter:
            mongo_query['created_date'] = date_filter
            
    docs = list(col.find(mongo_query).sort('created_date', -1)) if col is not None else []
    franchise_name_map = get_franchise_name_map()
    
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    wb = Workbook()
    ws = wb.active
    ws.title = "Material Requisitions"
    
    # Styles
    title_font = Font(name='Segoe UI', size=14, bold=True, color='1E293B')
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4B9EB0', end_color='4B9EB0', fill_type='solid')
    subtotal_font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
    subtotal_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    grand_total_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    grand_total_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    border_thin = Border(left=Side(style='thin', color='CBD5E1'),
                         right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'),
                         bottom=Side(style='thin', color='CBD5E1'))
                         
    # Title Row
    ws.append(["MATERIAL REQUISITIONS REPORT"])
    ws.cell(row=1, column=1).font = title_font
    ws.append([]) # empty row
    
    # Headers in bold and CAPS
    headers = ["DATE", "MR NUMBER", "FRANCHISE NAME", "STATUS", "ITEM ID", "MATERIAL / ITEM NAME", "QUANTITY", "UNIT RATE (₹)", "TOTAL AMOUNT (₹)"]
    ws.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if col_num in [1, 2, 4, 5] else ('right' if col_num in [7, 8, 9] else 'left'), vertical='center')
        
    grand_total_amount = 0.0
    current_row = 4
    
    for req in docs:
        fid = req.get('franchise_id') or ''
        fname = franchise_name_map.get(fid, fid)
        mr_num = req.get('mr_number', '')
        status_val = req.get('status', 'Draft')
        
        c_date = req.get('created_date')
        date_str = c_date.strftime('%d/%m/%Y') if isinstance(c_date, (datetime.datetime, datetime.date)) else str(c_date or '')[:10]
        
        raw_items = req.get('items') or []
        items = parse_items(raw_items)
        
        req_subtotal_amount = 0.0
        
        for idx, item in enumerate(items, start=1):
            if isinstance(item, dict):
                item_id = item.get('item_id', idx)
                item_name = item.get('item_name') or item.get('medicine_name') or ''
                qty = int(item.get('quantity', 1))
                rate = float(item.get('amount', 0))
                amount = qty * rate
                
                req_subtotal_amount += amount
                
                row_data = [
                    date_str if idx == 1 else "",
                    mr_num if idx == 1 else "",
                    fname if idx == 1 else "",
                    status_val if idx == 1 else "",
                    item_id,
                    item_name,
                    qty,
                    rate,
                    amount
                ]
                ws.append(row_data)
                for c_idx in range(1, 10):
                    c = ws.cell(row=current_row, column=c_idx)
                    c.border = border_thin
                    if c_idx in [1, 2, 4, 5]:
                        c.alignment = Alignment(horizontal='center', vertical='center')
                    elif c_idx in [7, 8, 9]:
                        c.alignment = Alignment(horizontal='right', vertical='center')
                        if c_idx in [8, 9]:
                            c.number_format = '#,##0.00'
                    else:
                        c.alignment = Alignment(horizontal='left', vertical='center')
                current_row += 1
                
        # Subtotal row for this MR - No quantity total, Total label in column 6, Amount in column 9 (Bold)
        subtotal_row = ["", "", "", "", "", f"Total for {mr_num}", "", "", req_subtotal_amount]
        ws.append(subtotal_row)
        for c_idx in range(1, 10):
            c = ws.cell(row=current_row, column=c_idx)
            c.font = subtotal_font
            c.fill = subtotal_fill
            c.border = border_thin
            if c_idx in [1, 2, 4, 5]:
                c.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx in [7, 8, 9]:
                c.alignment = Alignment(horizontal='right', vertical='center')
                if c_idx == 9:
                    c.number_format = '#,##0.00'
            else:
                c.alignment = Alignment(horizontal='left', vertical='center')
        current_row += 1
        
        grand_total_amount += req_subtotal_amount
        
    # Grand Total Row - Total in same column 9, Bold
    ws.append([]) # empty separator row
    current_row += 1
    grand_row = ["", "", "", "", "", "GRAND TOTAL", "", "", grand_total_amount]
    ws.append(grand_row)
    for c_idx in range(1, 10):
        c = ws.cell(row=current_row, column=c_idx)
        c.font = grand_total_font
        c.fill = grand_total_fill
        c.border = border_thin
        if c_idx in [1, 2, 4, 5]:
            c.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx in [7, 8, 9]:
            c.alignment = Alignment(horizontal='right', vertical='center')
            if c_idx == 9:
                c.number_format = '#,##0.00'
        else:
            c.alignment = Alignment(horizontal='left', vertical='center')
            
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Material_Requisitions_{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    wb.save(response)
    return response

def get_hms_collection(collection_name):
    try:
        client_dict = settings.DATABASES['default']['CLIENT']
        host = client_dict.get('host', 'mongodb://localhost:27017/')
        client = MongoClient(host)
        return client['HMS'][collection_name]
    except Exception as e:
        print("HMS Mongo connection error:", e)
        return None

@api_view(['GET'])
def get_material_items(request):
    try:
        department = request.query_params.get('department', 'DPT00002')
        item_col = get_hms_collection('hospital_itemmaster')
        grn_col = get_hms_collection('hospital_storesgrn')
        
        if item_col is None:
            return Response([], status=status.HTTP_200_OK)
            
        query = {'department': department}
        items_docs = list(item_col.find(query))
        
        result = []
        for it in items_docs:
            item_id = it.get('item_id')
            item_name = it.get('itemName') or ''
            
            tot_qty = 0
            try:
                tot_qty = int(it.get('total_quantity', 0) or 0)
            except Exception:
                pass
                
            app_qty = 0
            try:
                app_qty = int(it.get('approved_quantity', 0) or 0)
            except Exception:
                pass
                
            avail_qty = max(0, tot_qty - app_qty)
            
            # Lookup latest MRP from hospital_storesgrn
            mrp_val = 0.0
            if grn_col is not None and item_id:
                try:
                    grn_cursor = grn_col.find({'items.item_id': item_id}).sort('created_date', -1).limit(1)
                    grn_list = list(grn_cursor)
                    if grn_list:
                        for g_item in grn_list[0].get('items', []):
                            if g_item.get('item_id') == item_id:
                                raw_mrp = g_item.get('mrp') or g_item.get('unitPrice') or 0
                                try:
                                    mrp_val = float(str(raw_mrp))
                                except Exception:
                                    mrp_val = 0.0
                                break
                except Exception as ge:
                    print(f"Error fetching GRN for {item_id}:", ge)
                    
            result.append({
                "item_id": item_id,
                "item_name": item_name,
                "itemName": item_name,
                "department": it.get('department'),
                "group": it.get('group'),
                "category": it.get('category'),
                "total_quantity": tot_qty,
                "approved_quantity": app_qty,
                "available_quantity": avail_qty,
                "quantity": avail_qty,
                "amount": mrp_val,
                "mrp": mrp_val
            })
            
        return Response(result, status=status.HTTP_200_OK)
    except Exception as e:
        print("Error fetching material items:", e)
        return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

# Backward-compatible aliases
medicine_requisitions = material_requisitions
medicine_requisition_detail = material_requisition_detail


